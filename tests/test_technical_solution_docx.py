from pathlib import Path
import unittest

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "企业客服与工单执行Agent_MVP技术方案_v0.1.docx"


class TechnicalSolutionDocumentTest(unittest.TestCase):
    def test_exists_and_contains_required_sections(self):
        self.assertTrue(DOCX.exists())
        self.assertGreater(DOCX.stat().st_size, 20_000)

        doc = Document(DOCX)
        text = "\n".join(p.text for p in doc.paragraphs)
        required = [
            "企业客服与工单执行 Agent",
            "总体架构",
            "技术选型",
            "知识检索方案",
            "Elasticsearch 采用策略",
            "安全与审批边界",
            "可观测性与测试",
            "MVP 验收标准",
        ]
        for value in required:
            self.assertIn(value, text)

    def test_uses_expected_page_and_style_tokens(self):
        doc = Document(DOCX)
        section = doc.sections[0]
        self.assertEqual(round(section.page_width.inches, 2), 8.5)
        self.assertEqual(round(section.page_height.inches, 2), 11.0)
        self.assertEqual(round(section.top_margin.inches, 2), 1.0)
        self.assertEqual(round(section.left_margin.inches, 2), 1.0)

        normal = doc.styles["Normal"]
        self.assertEqual(normal.font.name, "Calibri")
        self.assertEqual(round(normal.font.size.pt, 1), 11.0)
        self.assertEqual(round(normal.paragraph_format.space_after.pt, 1), 6.0)

        heading1 = doc.styles["Heading 1"]
        self.assertEqual(round(heading1.font.size.pt, 1), 16.0)
        self.assertEqual(str(heading1.font.color.rgb), "2E74B5")
        self.assertEqual(round(heading1.paragraph_format.space_before.pt, 1), 16.0)
        self.assertEqual(round(heading1.paragraph_format.space_after.pt, 1), 8.0)

    def test_lists_and_table_geometry(self):
        doc = Document(DOCX)
        numbered_paragraphs = [
            p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None
        ]
        self.assertGreaterEqual(len(numbered_paragraphs), 35)
        self.assertGreaterEqual(len(doc.tables), 8)

        for table in doc.tables:
            tbl_pr = table._tbl.tblPr
            tbl_w = tbl_pr.find(qn("w:tblW"))
            tbl_ind = tbl_pr.find(qn("w:tblInd"))
            layout = tbl_pr.find(qn("w:tblLayout"))
            self.assertIsNotNone(tbl_w)
            self.assertEqual(tbl_w.get(qn("w:type")), "dxa")
            self.assertIsNotNone(tbl_ind)
            self.assertEqual(tbl_ind.get(qn("w:w")), "120")
            self.assertIsNotNone(layout)
            self.assertEqual(layout.get(qn("w:type")), "fixed")
            grid_width = sum(int(col.get(qn("w:w"))) for col in table._tbl.tblGrid)
            self.assertEqual(grid_width, 9360)

    def test_no_placeholder_or_internal_citation_tokens(self):
        doc = Document(DOCX)
        all_text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        self.assertNotIn("TODO", all_text)
        self.assertNotIn("turn0search", all_text)
        self.assertNotIn(":codex-file-citation", all_text)


if __name__ == "__main__":
    unittest.main()
