from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "企业客服与工单执行Agent_MVP技术方案_v0.1.docx"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
BORDER = "D0D5DD"
WHITE = "FFFFFF"
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT_LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            set_cell_border(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_numbering(doc, num_format, text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_list_item(doc, text, num_id):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    apply_numbering(paragraph, num_id)
    set_font(paragraph.add_run(text), size=11)
    return paragraph


def add_paragraph(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_font(paragraph.add_run(bold_prefix), size=11, bold=True, color=NAVY)
        set_font(paragraph.add_run(text[len(bold_prefix):]), size=11)
    else:
        set_font(paragraph.add_run(text), size=11)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    for run in paragraph.runs:
        set_font(run, bold=True)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_callout(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(10)
    paragraph.paragraph_format.right_indent = Pt(10)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.10
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    ppr.append(borders)
    set_font(paragraph.add_run(f"{label}  "), size=11, bold=True, color=NAVY)
    set_font(paragraph.add_run(text), size=11, color=NAVY)


def add_table(doc, headers, rows, widths_dxa, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_fill(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        set_font(paragraph.add_run(header), size=9.5, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            set_font(paragraph.add_run(value), size=9.5)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    set_font(field_run, size=9, color=MUTED)
    field_run._r.append(fld_char1)
    field_run._r.append(instr)
    field_run._r.append(fld_char2)
    end_run = paragraph.add_run(" 页")
    set_font(end_run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run("HARBOR SUPPORT  ·  MVP 技术方案"), size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    add_page_field(footer)


def add_masthead(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_font(kicker.add_run("TECHNICAL SOLUTION"), size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("企业客服与工单执行 Agent"), size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("MVP 技术方案 · v0.1"), size=14, color=MUTED)

    metadata = [
        ("文档状态", "建议方案"),
        ("适用阶段", "核心流程验证与首版工程实现"),
        ("编制日期", "2026-08-31"),
        ("核心决策", "模块化单体 + 单 Agent + PostgreSQL/pgvector"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0
        set_font(paragraph.add_run(f"{label}："), size=10.5, color=NAVY, bold=True)
        set_font(paragraph.add_run(value), size=10.5, color="344054")

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "推荐结论",
        "首版采用 Next.js + FastAPI + OpenAI Agents SDK + PostgreSQL/pgvector。"
        "Elasticsearch 暂不作为 MVP 基础依赖，在知识规模、中文关键词检索或跨域搜索出现明确瓶颈后再引入。",
    )


def build_document():
    doc = Document()
    configure_document(doc)
    bullet_num_id = add_numbering(doc, "bullet", "•")
    decimal_num_id = add_numbering(doc, "decimal", "%1.")
    add_masthead(doc)

    add_heading(doc, "1. 方案背景与目标", 1)
    add_paragraph(
        doc,
        "本方案基于《企业客服与工单执行 Agent MVP 产品设计文档 v0.1》和现有前端 Demo，"
        "目标是在不过度建设基础设施的前提下，跑通政策问答、订单与物流查询、异常建单、退款确认和结果回写。",
    )
    add_list_item(doc, "业务事实必须来自知识库或后端工具，模型不得编造订单、物流、工单或退款结果。", bullet_num_id)
    add_list_item(doc, "写操作必须经过后端规则校验；退款等高风险动作必须获得用户明确确认。", bullet_num_id)
    add_list_item(doc, "对话、工单、退款与工具调用需要持久化和可审计，并支持重复演示与自动化测试。", bullet_num_id)
    add_list_item(doc, "MVP 以模块化单体为主，避免过早引入分布式队列、多 Agent 和复杂工作流引擎。", bullet_num_id)

    add_heading(doc, "2. 总体架构", 1)
    add_paragraph(
        doc,
        "系统采用前后端分离的模块化单体。前端负责对话、业务上下文和确认交互；FastAPI 提供业务 API 与 Agent 运行入口；"
        "Agent 只负责理解诉求、选择工具和组织回答；领域服务负责权限、规则、事务、幂等与状态转换；PostgreSQL 是唯一业务事实源。",
    )
    architecture_rows = [
        ("交互层", "Next.js Web", "客服对话、来源展示、工单详情、退款确认、错误与执行状态"),
        ("接口层", "FastAPI", "鉴权、请求校验、流式响应、Agent 入口、业务 API"),
        ("编排层", "OpenAI Agents SDK", "单 Agent、工具选择、结构化输出、guardrails、会话与 tracing"),
        ("领域层", "Application Services", "订单、物流、工单、退款、知识、审计与幂等规则"),
        ("数据层", "PostgreSQL + pgvector", "事务数据、会话消息、知识分块、向量、工具调用日志"),
        ("外部层", "模拟/真实业务适配器", "电商、物流、支付、通知等接口；MVP 首先使用模拟实现"),
    ]
    add_table(doc, ["层级", "建议组件", "职责"], architecture_rows, [1550, 2500, 5310])
    add_callout(
        doc,
        "架构边界",
        "LLM 产生的是意图和工具参数建议，不直接修改业务数据。所有工具调用必须进入领域服务，"
        "由确定性代码完成授权、金额计算、状态校验、事务提交和审计记录。",
    )

    add_heading(doc, "3. 技术选型", 1)
    stack_rows = [
        ("前端", "Next.js + React + TypeScript", "沿用现有 React 界面；正式版使用稳定 Next.js，便于 Docker 自托管"),
        ("样式与组件", "Tailwind CSS + shadcn/ui", "复用 Demo 的视觉与组件资产，减少重做"),
        ("后端", "FastAPI + Pydantic", "类型明确、异步友好，便于把工具输入输出建模为结构化契约"),
        ("Agent", "OpenAI Agents SDK", "单 Agent、function tools、guardrails、会话与 tracing"),
        ("数据访问", "SQLAlchemy 2 + Alembic", "事务、迁移和可测试的数据访问层"),
        ("主数据库", "PostgreSQL", "订单、工单、退款、对话和审计的唯一事实源"),
        ("知识检索", "pgvector + 元数据过滤", "MVP 先满足语义检索与可追溯来源；保留混合检索扩展点"),
        ("测试", "pytest + Playwright", "覆盖领域规则、工具契约、Agent 流程和端到端交互"),
        ("部署", "Docker Compose", "前端、API、PostgreSQL 一键启动；后续再拆分基础设施"),
        ("可观测性", "结构化日志 + OpenTelemetry", "串联 conversation_id、ticket_id、tool_call_id 和 trace_id"),
    ]
    add_table(doc, ["领域", "技术", "选择理由"], stack_rows, [1500, 2850, 5010])
    add_paragraph(
        doc,
        "现有 Demo 中的 React、Tailwind、shadcn/ui 和 Lucide 可继续使用。当前 vinext beta 与 Cloudflare 专用配置适合演示站，"
        "正式 MVP 建议切换到稳定 Next.js 构建，以匹配产品文档中的 Docker Compose 交付方式。",
    )

    add_heading(doc, "4. 核心模块设计", 1)
    add_heading(doc, "4.1 Web 客户端", 2)
    add_list_item(doc, "对话区展示用户消息、Agent 回复、知识来源、工具执行状态与可恢复错误。", bullet_num_id)
    add_list_item(doc, "业务上下文区展示订单、物流、工单和退款状态，但只展示当前用户有权访问的数据。", bullet_num_id)
    add_list_item(doc, "退款确认卡片必须明确金额、方式、原因和影响；确认与取消是独立、可审计的用户动作。", bullet_num_id)
    add_list_item(doc, "Agent 响应建议使用流式传输；业务工具结果以结构化事件更新 UI，而非从自然语言中反向解析。", bullet_num_id)

    add_heading(doc, "4.2 Agent 编排与工具", 2)
    add_paragraph(
        doc,
        "MVP 使用一个 Customer Support Agent。系统指令规定能力边界和回答格式，工具使用 Pydantic 模型声明输入输出；"
        "涉及业务写入的工具需要审批状态和服务端 guardrail。",
    )
    tools_rows = [
        ("search_knowledge_base", "读", "查询政策分块并返回文章、版本、条款和相关度"),
        ("get_order", "读", "校验订单归属后返回订单状态、实付与可退金额"),
        ("get_shipping_status", "读", "返回物流节点并由确定性规则判断异常"),
        ("create_ticket", "写", "幂等创建工单，关联订单、对话和证据"),
        ("get_ticket", "读", "返回工单状态和处理记录"),
        ("create_refund_request", "写", "创建 WAITING_APPROVAL 申请，不执行资金动作"),
        ("confirm_refund", "高风险写", "校验用户确认、金额、状态和幂等键后模拟执行退款"),
    ]
    add_table(doc, ["工具", "类型", "服务端约束"], tools_rows, [2600, 1400, 5360])

    add_heading(doc, "4.3 领域服务", 2)
    add_list_item(doc, "Identity & Authorization：解析当前用户，执行订单归属和资源访问校验。", bullet_num_id)
    add_list_item(doc, "Order & Shipping：封装订单与物流适配器，统一超时、重试和错误语义。", bullet_num_id)
    add_list_item(doc, "Ticket：维护 OPEN、WAITING_APPROVAL、RESOLVED 状态及合法转换。", bullet_num_id)
    add_list_item(doc, "Refund：计算可退金额，生成退款申请，处理确认、幂等和结果回写。", bullet_num_id)
    add_list_item(doc, "Audit：在同一事务或可靠 outbox 中记录工具输入摘要、输出摘要、耗时和结果。", bullet_num_id)

    add_heading(doc, "5. 数据与一致性设计", 1)
    add_paragraph(
        doc,
        "PostgreSQL 保存 customers、orders、knowledge_articles、knowledge_chunks、conversations、messages、tickets、"
        "refund_requests、tool_invocations 和 idempotency_records。所有表使用 UUID 主键、created_at/updated_at，关键状态变更记录版本号。",
    )
    add_list_item(doc, "工单与退款状态通过数据库事务更新，不从对话文本推断最终状态。", bullet_num_id)
    add_list_item(doc, "退款确认请求必须携带 idempotency_key；重复请求返回首次执行结果。", bullet_num_id)
    add_list_item(doc, "工具调用记录保存脱敏后的参数、结果、耗时、错误类型、模型运行和关联业务对象。", bullet_num_id)
    add_list_item(doc, "知识文章和分块保存版本、有效期、来源 URI 与内容哈希，便于回答引用和增量重建向量。", bullet_num_id)
    add_list_item(doc, "如后续接入异步任务，优先使用数据库 outbox，避免请求同时写 PostgreSQL 和搜索系统。", bullet_num_id)

    add_heading(doc, "6. 知识检索方案", 1)
    add_paragraph(
        doc,
        "MVP 知识量较小，采用 pgvector 进行向量召回，并结合文章类型、版本、有效期和渠道等元数据过滤。"
        "召回结果必须包含来源和条款，低于阈值时 Agent 明确表示无法确认并建议人工跟进。",
    )
    add_list_item(doc, "入库：规范化文本，按条款或主题分块，生成 embedding，保存内容哈希和版本。", decimal_num_id)
    add_list_item(doc, "召回：向量 Top-K + 元数据过滤；必要时补充轻量关键词匹配。", decimal_num_id)
    add_list_item(doc, "重排：根据相关度、版本有效性和业务优先级排序，返回引用片段。", decimal_num_id)
    add_list_item(doc, "回答：只基于召回证据生成结论，并把文章名、版本和条款返回给前端。", decimal_num_id)

    add_heading(doc, "6.1 Elasticsearch 采用策略", 2)
    add_callout(
        doc,
        "当前决策",
        "MVP 不引入 Elasticsearch。它适合作为后续的搜索读模型，而不是订单、工单、退款或会话的主数据库。",
    )
    es_rows = [
        ("继续使用 PostgreSQL/pgvector", "知识规模小、检索 QPS 低、以语义问答为主、团队希望保持单一数据系统"),
        ("评估引入 Elasticsearch", "需要精细中文分词、同义词/短语/高亮、复杂字段权重、跨知识与工单搜索，或性能指标出现稳定瓶颈"),
        ("引入后的数据路径", "PostgreSQL 仍为事实源；通过 outbox/CDC 异步构建 Elasticsearch 索引，禁止业务请求双写"),
    ]
    add_table(doc, ["决策点", "判断条件"], es_rows, [2600, 6760])

    add_heading(doc, "7. 关键业务流程", 1)
    add_heading(doc, "7.1 政策问答", 2)
    add_list_item(doc, "识别政策诉求并提取必要上下文。", decimal_num_id)
    add_list_item(doc, "检索有效知识分块并返回来源；无足够证据时停止确定性回答。", decimal_num_id)
    add_list_item(doc, "Agent 组织面向用户的答案，前端展示来源和版本。", decimal_num_id)

    add_heading(doc, "7.2 物流异常与建单", 2)
    add_list_item(doc, "校验当前用户与订单归属，失败时不返回订单信息。", decimal_num_id)
    add_list_item(doc, "调用物流适配器；由确定性规则判断是否达到异常阈值。", decimal_num_id)
    add_list_item(doc, "用户同意或规则允许后创建 SHIPPING 工单，并返回编号与预计处理时间。", decimal_num_id)

    add_heading(doc, "7.3 退款确认闭环", 2)
    add_list_item(doc, "查询订单和政策，服务端计算可退金额。", decimal_num_id)
    add_list_item(doc, "创建 WAITING_APPROVAL 退款申请，前端展示金额、方式和原因。", decimal_num_id)
    add_list_item(doc, "用户明确确认后提交 idempotency_key；服务端再次校验归属、金额与状态。", decimal_num_id)
    add_list_item(doc, "模拟执行退款，原子更新申请与工单为 RESOLVED，并记录审计事件。", decimal_num_id)

    add_heading(doc, "8. 安全与审批边界", 1)
    security_rows = [
        ("身份与越权", "所有订单、工单和退款查询先做服务端归属校验；禁止模型传入的 user_id 成为可信依据"),
        ("高风险操作", "退款必须经过明确确认；金额、原因、状态和幂等键由服务端校验"),
        ("Prompt Injection", "知识内容仅作为不可信数据；工具权限由代码和 allowlist 决定，检索文本不能改变系统规则"),
        ("隐私与日志", "日志默认脱敏，不记录完整支付信息、密钥或无必要的个人信息"),
        ("工具错误", "超时、限流、业务拒绝和系统错误使用明确错误码；Agent 不得把失败包装为成功"),
        ("审计", "记录操作主体、确认时间、工具调用、业务对象、结果与 trace_id"),
    ]
    add_table(doc, ["风险", "控制措施"], security_rows, [2000, 7360])

    add_heading(doc, "9. 可观测性与测试", 1)
    add_heading(doc, "9.1 可观测性", 2)
    add_list_item(doc, "统一关联 conversation_id、message_id、tool_call_id、ticket_id、refund_request_id 和 trace_id。", bullet_num_id)
    add_list_item(doc, "记录 Agent 耗时、模型用量、工具成功率、知识命中率、确认转化率和重复请求命中率。", bullet_num_id)
    add_list_item(doc, "模型输入输出只在受控环境中按需保留，并支持关闭敏感内容 tracing。", bullet_num_id)

    add_heading(doc, "9.2 测试策略", 2)
    test_rows = [
        ("单元测试", "领域状态机、金额计算、订单归属、幂等、知识阈值和异常分类"),
        ("工具契约测试", "每个 function tool 的 schema、授权、成功和失败返回"),
        ("集成测试", "FastAPI + PostgreSQL 事务、迁移、outbox 和模拟外部适配器"),
        ("Agent 场景测试", "四个核心场景、知识不足、越权、工具失败、重复确认和拒绝退款"),
        ("端到端测试", "Playwright 覆盖对话、建单、退款确认、刷新后状态恢复"),
    ]
    add_table(doc, ["层级", "重点"], test_rows, [2000, 7360])
    add_paragraph(
        doc,
        "首版至少提供 20 个稳定自动化用例。涉及模型的测试使用固定工具桩和结构化断言，避免只比较自然语言全文。",
    )

    add_heading(doc, "10. 部署与环境", 1)
    add_paragraph(
        doc,
        "本地与演示环境通过 Docker Compose 启动 web、api 和 postgres 三个核心服务。"
        "前端与 API 使用环境变量配置模型、数据库和外部适配器；密钥不进入镜像、仓库或工具日志。",
    )
    add_list_item(doc, "开发环境：本地热更新，PostgreSQL 使用持久化卷，模拟订单/物流/支付适配器。", bullet_num_id)
    add_list_item(doc, "测试环境：独立数据库，执行迁移、pytest 与 Playwright，测试结束后清理数据。", bullet_num_id)
    add_list_item(doc, "演示环境：固定种子数据和可重置场景，所有外部写操作仍为模拟执行。", bullet_num_id)
    add_list_item(doc, "生产演进：优先采用托管 PostgreSQL、对象存储和集中日志；达到明确容量瓶颈后再拆服务。", bullet_num_id)

    add_heading(doc, "11. 分阶段实施建议", 1)
    phases = [
        ("阶段 1 · 工程骨架", "Next.js/FastAPI/PostgreSQL、迁移、鉴权、种子数据、统一错误与 trace_id"),
        ("阶段 2 · 核心闭环", "知识检索、订单物流工具、工单状态机、退款确认与幂等"),
        ("阶段 3 · 质量交付", "20+ 自动化测试、Docker Compose、演示脚本、日志与指标"),
        ("阶段 4 · 搜索增强", "基于真实检索评测决定是否增加关键词重排或 Elasticsearch"),
    ]
    add_table(doc, ["阶段", "范围"], phases, [2400, 6960])

    add_heading(doc, "12. MVP 验收标准", 1)
    add_list_item(doc, "四个核心场景能够重复执行，结果来自工具或知识来源，过程可追踪。", bullet_num_id)
    add_list_item(doc, "越权订单不泄露信息；知识不足不编造答案；工具失败不伪装为成功。", bullet_num_id)
    add_list_item(doc, "退款确认前不执行；重复确认不会生成第二次退款。", bullet_num_id)
    add_list_item(doc, "对话、工单、退款、工具调用和审计事件均持久化。", bullet_num_id)
    add_list_item(doc, "自动化测试、文档和 Docker Compose 验证全部通过。", bullet_num_id)

    add_heading(doc, "13. 主要风险与应对", 1)
    risk_rows = [
        ("模型误选工具", "工具 allowlist、结构化参数、服务端 guardrail 和场景评测"),
        ("业务规则散落在 Prompt", "规则统一放在领域服务和数据库约束，Prompt 只描述交互边界"),
        ("中文检索质量不足", "建立小型评测集，先调分块/embedding/阈值，再决定是否引入 Elasticsearch"),
        ("外部接口不稳定", "适配器超时、有限重试、错误分类和人工工单兜底"),
        ("过早复杂化", "以可测业务闭环为里程碑，暂缓 Redis、Celery、Kafka、多 Agent 与 Kubernetes"),
    ]
    add_table(doc, ["风险", "应对"], risk_rows, [2400, 6960])

    add_heading(doc, "14. 参考资料", 1)
    references = [
        "企业客服与工单执行 Agent MVP 产品设计文档 v0.1（本项目 docs 目录）",
        "OpenAI Agents SDK Documentation: https://openai.github.io/openai-agents-python/",
        "pgvector Documentation: https://github.com/pgvector/pgvector",
        "Elasticsearch Hybrid Search: https://www.elastic.co/docs/solutions/search/hybrid-search",
        "Next.js Self-Hosting: https://nextjs.org/docs/app/guides/self-hosting",
    ]
    for ref in references:
        add_list_item(doc, ref, bullet_num_id)

    doc.core_properties.title = "企业客服与工单执行 Agent MVP 技术方案"
    doc.core_properties.subject = "MVP 技术架构与实施建议"
    doc.core_properties.author = "Harbor Support Project"
    doc.core_properties.keywords = "Agent, FastAPI, Next.js, PostgreSQL, pgvector, Elasticsearch"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
