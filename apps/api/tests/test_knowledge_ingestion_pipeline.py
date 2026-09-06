import base64
import io
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select

from serviceops.knowledge.ingestion import (
    CHUNKING_VERSION,
    MAX_DOCUMENT_BYTES,
    PARSER_VERSION,
    KnowledgePipelineError,
    chunk_blocks,
    ingest_document,
    parse_document,
)
from serviceops.knowledge.releases import (
    KnowledgeReleaseError,
    approve_knowledge_release,
    create_knowledge_release,
    evaluate_knowledge_release,
    publish_knowledge_release,
    rollback_knowledge_release,
)
from serviceops.models import (
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeRelease,
)
from serviceops.seed import OPS_SESSION_TOKEN

OPS_HEADERS = {"X-Ops-Session": OPS_SESSION_TOKEN}


def _pdf_fixture(text: str = "七天退货规则") -> bytes:
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


def _docx_fixture() -> bytes:
    document = Document()
    document.add_heading("售后规则", level=1)
    document.add_paragraph("签收后 7 天内可以申请无理由退货。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "场景"
    table.rows[0].cells[1].text = "处理"
    table.add_row().cells[0].text = "质量问题"
    table.rows[-1].cells[1].text = "上传凭证"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_parsers_emit_structured_traceable_blocks_and_chunks():
    markdown = """# 售后规则

## 退货

签收后 7 天内可以申请无理由退货。

- 商品完整
- 赠品完整

| 场景 | 时限 |
| --- | --- |
| 质量问题 | 15 天 |
"""
    parsed = parse_document("rules.md", markdown.encode())
    chunks = chunk_blocks(parsed.blocks)

    assert parsed.parser_version == PARSER_VERSION
    assert parsed.page_count == 1
    assert {block.kind for block in parsed.blocks} >= {"heading", "paragraph", "list", "table"}
    assert any(block.title_path == ("售后规则", "退货") for block in parsed.blocks)
    assert all(block.start_offset < block.end_offset for block in parsed.blocks)
    assert chunks
    assert all(chunk.content_hash and chunk.end_offset > chunk.start_offset for chunk in chunks)
    assert all(len(chunk.content) <= 1200 for chunk in chunks)


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("rules.txt", "第 2.1 条\n签收后 7 天内可以退货。".encode()),
        ("rules.pdf", _pdf_fixture()),
        ("rules.docx", _docx_fixture()),
    ],
    ids=["txt", "pdf", "docx"],
)
def test_supported_parsers_share_one_output_contract(filename, content):
    parsed = parse_document(filename, content)

    assert parsed.filename == filename
    assert parsed.media_type
    assert parsed.content_hash
    assert parsed.blocks
    assert all(block.text.strip() for block in parsed.blocks)


@pytest.mark.parametrize(
    ("filename", "content", "code"),
    [
        ("rules.exe", b"not supported", "DOCUMENT_TYPE_UNSUPPORTED"),
        ("rules.txt", b"   \n", "DOCUMENT_EMPTY"),
        ("rules.pdf", b"not a pdf", "DOCUMENT_PARSE_FAILED"),
        ("rules.pdf", b"%PDF-1.4\n/Encrypt true\n%%EOF", "DOCUMENT_ENCRYPTED"),
    ],
)
def test_parser_failures_close_without_partial_output(filename, content, code):
    with pytest.raises(KnowledgePipelineError) as error:
        parse_document(filename, content)
    assert error.value.code == code

    with pytest.raises(KnowledgePipelineError) as too_large:
        parse_document("rules.txt", b"x" * (MAX_DOCUMENT_BYTES + 1))
    assert too_large.value.code == "DOCUMENT_TOO_LARGE"


def test_ingestion_is_idempotent_and_updates_create_new_document(db):
    first = ingest_document(
        db,
        filename="rules.md",
        source_uri="fixture://stage2/rules.md",
        data=b"# V1\n\n7 days.",
        idempotency_key="stage2-upload-1",
    )
    second = ingest_document(
        db,
        filename="rules.md",
        source_uri="fixture://stage2/rules.md",
        data=b"# V1\n\n7 days.",
        idempotency_key="stage2-upload-1",
    )
    updated = ingest_document(
        db,
        filename="rules.md",
        source_uri="fixture://stage2/rules.md",
        data=b"# V2\n\n15 days.",
        idempotency_key="stage2-upload-2",
    )

    assert first.document.id == second.document.id
    assert second.duplicate is True
    assert updated.document.id != first.document.id
    assert updated.document.supersedes_document_id == first.document.id
    assert updated.document.chunking_version == CHUNKING_VERSION
    assert db.scalar(select(KnowledgeDocument).where(KnowledgeDocument.id == first.document.id))
    assert db.scalar(
        select(KnowledgeChunk.id).where(KnowledgeChunk.document_id == first.document.id)
    )
    assert len(list(db.scalars(select(KnowledgeDocument)))) == 2


def test_ingestion_rejects_external_sources_and_conflicting_idempotency(db):
    with pytest.raises(KnowledgePipelineError) as external:
        ingest_document(
            db,
            filename="rules.txt",
            source_uri="https://example.invalid/rules.txt",
            data=b"rules",
            idempotency_key="external-source",
        )
    assert external.value.code == "KNOWLEDGE_EXTERNAL_SOURCE_FORBIDDEN"

    ingest_document(
        db,
        filename="rules.txt",
        source_uri="fixture://stage2/idempotent.txt",
        data=b"first",
        idempotency_key="same-key",
    )
    with pytest.raises(Exception) as conflict:
        ingest_document(
            db,
            filename="rules.txt",
            source_uri="fixture://stage2/idempotent.txt",
            data=b"second",
            idempotency_key="same-key",
        )
    assert getattr(conflict.value, "code", None) == "KNOWLEDGE_IDEMPOTENCY_CONFLICT"


def test_release_evaluation_approval_publish_and_rollback_are_snapshot_based(db):
    first_document = ingest_document(
        db,
        filename="v1.md",
        source_uri="fixture://stage2/v1.md",
        data=b"# V1\n\n7 days.",
    ).document
    first = create_knowledge_release(
        db,
        release_version="stage2-v1",
        document_ids=[first_document.id],
        created_by="tester",
        git_commit="test-commit-1",
    )
    first = evaluate_knowledge_release(db, first.id)
    assert first.status == "EVALUATED"
    assert first.evaluation_report["passed"] is True
    first = approve_knowledge_release(db, first.id, approved_by="reviewer")
    first = publish_knowledge_release(db, first.id)
    assert first.status == "PUBLISHED"

    second_document = ingest_document(
        db,
        filename="v2.md",
        source_uri="fixture://stage2/v1.md",
        data=b"# V2\n\n15 days.",
    ).document
    second = create_knowledge_release(
        db,
        release_version="stage2-v2",
        document_ids=[second_document.id],
        created_by="tester",
        git_commit="test-commit-2",
    )
    second = evaluate_knowledge_release(db, second.id)
    second = approve_knowledge_release(db, second.id, approved_by="reviewer")
    second = publish_knowledge_release(db, second.id)
    assert second.status == "PUBLISHED"
    assert db.get(KnowledgeRelease, first.id).status == "RETIRED"

    rolled_back = rollback_knowledge_release(db, second.id, target_release_id=first.id)
    assert rolled_back.id == first.id
    assert rolled_back.status == "PUBLISHED"
    assert db.get(KnowledgeRelease, second.id).status == "ROLLED_BACK"
    assert db.get(KnowledgeRelease, second.id).rollback_release_id == first.id


def test_failed_release_evaluation_cannot_be_approved(db, monkeypatch):
    document = ingest_document(
        db,
        filename="failed-evaluation.md",
        source_uri="fixture://stage2/failed-evaluation.md",
        data=b"# V1\n\n7 days.",
    ).document
    release = create_knowledge_release(
        db,
        release_version="stage2-failed-evaluation",
        document_ids=[document.id],
        created_by="tester",
    )
    monkeypatch.setattr(
        "serviceops.knowledge.releases.evaluate_knowledge_search",
        lambda _db: {"dataset_version": "knowledge-baseline-30-v1", "passed": False, "metrics": {}},
    )
    release = evaluate_knowledge_release(db, release.id)
    assert release.status == "REJECTED"
    with pytest.raises(KnowledgeReleaseError) as error:
        approve_knowledge_release(db, release.id, approved_by="reviewer")
    assert error.value.code == "KNOWLEDGE_RELEASE_EVALUATION_FAILED"


def test_knowledge_pipeline_api_exposes_preview_and_lifecycle(client):
    response = client.post("/api/ops/knowledge/documents", headers=OPS_HEADERS, json={
        "filename": "api-rules.md",
        "source_uri": "fixture://stage2/api-rules.md",
        "content_base64": base64.b64encode(b"# API\n\n7 days.").decode(),
        "idempotency_key": "api-stage2-1",
    })
    assert response.status_code == 200
    document = response.json()
    assert document["status"] == "PARSED"

    chunks_response = client.get(
        f"/api/ops/knowledge/documents/{document['id']}/chunks",
        headers=OPS_HEADERS,
    )
    assert chunks_response.status_code == 200
    assert chunks_response.json()[0]["source_uri"] == "fixture://stage2/api-rules.md"

    release_response = client.post(
        "/api/ops/knowledge/releases",
        headers=OPS_HEADERS,
        json={
            "release_version": "api-stage2-v1",
            "document_ids": [document["id"]],
            "git_commit": "api-test",
        },
    )
    assert release_response.status_code == 200
    release_id = release_response.json()["id"]
    assert release_response.json()["status"] == "DRAFT"
    assert client.post(
        f"/api/ops/knowledge/releases/{release_id}/evaluate", headers=OPS_HEADERS
    ).json()["status"] == "EVALUATED"
    assert client.post(
        f"/api/ops/knowledge/releases/{release_id}/approve", headers=OPS_HEADERS
    ).json()["status"] == "APPROVED"
    assert client.post(
        f"/api/ops/knowledge/releases/{release_id}/publish", headers=OPS_HEADERS
    ).json()["status"] == "PUBLISHED"
    assert client.get(
        f"/api/ops/knowledge/releases/{release_id}/chunks", headers=OPS_HEADERS
    ).status_code == 200


def test_knowledge_pipeline_api_preserves_role_boundary(client):
    body = {
        "filename": "api-rules.md",
        "source_uri": "fixture://stage2/forbidden.md",
        "content_base64": base64.b64encode(b"rules").decode(),
    }
    assert client.post("/api/ops/knowledge/documents", json=body).status_code == 403
    response = client.post("/api/ops/knowledge/documents", headers=OPS_HEADERS, json={**body, "source_uri": "https://example.invalid/rules.md"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "KNOWLEDGE_EXTERNAL_SOURCE_FORBIDDEN"


def test_stage2_code_does_not_enable_embedding_or_external_network_calls():
    root = Path(__file__).parents[3]
    source = "\n".join(
        (root / path).read_text(encoding="utf-8")
        for path in (
            "apps/api/src/serviceops/knowledge/ingestion.py",
            "apps/api/src/serviceops/knowledge/releases.py",
        )
    )
    assert "fetch(" not in source
    assert "requests." not in source
    assert "WebSocket" not in source
    assert "embedding" not in source.lower()
    assert "rrf" not in source.lower()
